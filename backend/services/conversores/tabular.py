"""Leitura de dados estruturados (CSV, JSON, YAML, XML, XLS, XLSX, DOCX)."""

import csv
import io
import json
import re
from datetime import date, datetime
from typing import Any

import yaml
from defusedxml import ElementTree as DefusedET
from defusedxml.common import DefusedXmlException
from docx import Document
from openpyxl import load_workbook

from .erros import ConversaoError

MAX_LINHAS = 50_000
_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

Linhas = list[list[Any]]


def limpar(valor: Any) -> Any:
    """Remove caracteres de controle que formatos XML (docx/xlsx) não aceitam."""
    return _CONTROLE.sub("", valor) if isinstance(valor, str) else valor


def decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ConversaoError("Não foi possível identificar a codificação do arquivo.")


def _celula(valor: Any) -> Any:
    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)
    return "" if valor is None else valor


def _limite(total: int) -> None:
    if total > MAX_LINHAS:
        raise ConversaoError("O arquivo excede o limite de linhas permitido.", 413)


def _linhas_de_dados(dados: Any) -> Linhas:
    if isinstance(dados, dict):
        if len(dados) == 1 and isinstance(next(iter(dados.values())), list):
            dados = next(iter(dados.values()))
        else:
            dados = [dados]
    if not isinstance(dados, list):
        dados = [{"valor": dados}]
    _limite(len(dados))
    if not dados:
        return []

    if all(isinstance(item, dict) for item in dados):
        colunas: list[str] = []
        for item in dados:
            for chave in item:
                if chave not in colunas:
                    colunas.append(chave)
        return [colunas] + [[_celula(item.get(c)) for c in colunas] for item in dados]
    if all(isinstance(item, list) for item in dados):
        return [[_celula(v) for v in item] for item in dados]
    return [["valor"]] + [[_celula(item)] for item in dados]


def linhas_para_registros(linhas: Linhas) -> list[dict[str, Any]]:
    if not linhas:
        return []
    cabecalho, *corpo = linhas
    nomes = [str(c).strip() or f"coluna_{i + 1}" for i, c in enumerate(cabecalho)]
    return [
        {nome: (linha[i] if i < len(linha) else "") for i, nome in enumerate(nomes)}
        for linha in corpo
    ]


def ler_csv(conteudo: bytes) -> Linhas:
    texto = decodificar(conteudo)
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialeto = csv.excel
    linhas: Linhas = []
    for linha in csv.reader(io.StringIO(texto, newline=""), dialeto):
        linhas.append(linha)
        _limite(len(linhas))
    return linhas


def carregar_json(conteudo: bytes) -> Any:
    try:
        return json.loads(decodificar(conteudo))
    except json.JSONDecodeError as erro:
        raise ConversaoError(f"JSON inválido: {erro.msg} (linha {erro.lineno}).") from erro


def carregar_yaml(conteudo: bytes) -> Any:
    texto = decodificar(conteudo)
    try:
        # Aliases podem expandir exponencialmente ao serializar; não são aceitos.
        if any(isinstance(e, yaml.AliasEvent) for e in yaml.parse(texto, Loader=yaml.SafeLoader)):
            raise ConversaoError("YAML com âncoras e aliases (&, *) não é suportado.")
        return yaml.safe_load(texto)
    except yaml.YAMLError as erro:
        raise ConversaoError("YAML inválido ou com mais de um documento.") from erro


def _elemento(el) -> Any:
    dados: dict[str, Any] = {f"@{k}": v for k, v in el.attrib.items()}
    filhos = list(el)
    for filho in filhos:
        valor = _elemento(filho)
        if filho.tag in dados:
            atual = dados[filho.tag]
            if isinstance(atual, list):
                atual.append(valor)
            else:
                dados[filho.tag] = [atual, valor]
        else:
            dados[filho.tag] = valor
    texto = (el.text or "").strip()
    if not filhos and not el.attrib:
        return texto
    if texto:
        dados["#text"] = texto
    return dados


def carregar_xml(conteudo: bytes) -> Any:
    try:
        raiz = DefusedET.fromstring(conteudo)
        return {raiz.tag: _elemento(raiz)}
    except DefusedXmlException as erro:
        raise ConversaoError("XML com entidades ou DTD externos não é permitido.") from erro
    except (DefusedET.ParseError, RecursionError) as erro:
        raise ConversaoError("XML inválido ou profundo demais.") from erro


def _registros_xml(conteudo: bytes) -> Any:
    tag, raiz = next(iter(carregar_xml(conteudo).items()))
    if not isinstance(raiz, dict):
        return [{tag: raiz}]
    valores = list(raiz.values())
    if len(valores) == 1 and isinstance(valores[0], list):
        return valores[0]
    return [raiz]


def ler_json(conteudo: bytes) -> Linhas:
    return _linhas_de_dados(carregar_json(conteudo))


def ler_yaml(conteudo: bytes) -> Linhas:
    return _linhas_de_dados(carregar_yaml(conteudo))


def ler_xml(conteudo: bytes) -> Linhas:
    return _linhas_de_dados(_registros_xml(conteudo))


def ler_xlsx(conteudo: bytes) -> Linhas:
    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        raise ConversaoError("Planilha Excel inválida ou corrompida.") from erro
    try:
        aba = planilha.worksheets[0]
        linhas: Linhas = []
        for linha in aba.iter_rows(values_only=True):
            linhas.append([_celula(v) for v in linha])
            _limite(len(linhas))
    finally:
        planilha.close()
    while linhas and all(v == "" for v in linhas[-1]):
        linhas.pop()
    return linhas


def ler_xls(conteudo: bytes) -> Linhas:
    import xlrd

    try:
        livro = xlrd.open_workbook(file_contents=conteudo, on_demand=True)
        aba = livro.sheet_by_index(0)
    except Exception as erro:
        raise ConversaoError("Planilha .xls inválida, protegida ou corrompida.") from erro

    def valor(celula):
        if celula.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(celula.value, livro.datemode).isoformat(sep=" ")
            except Exception:
                return celula.value
        if celula.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(celula.value)
        if celula.ctype == xlrd.XL_CELL_NUMBER and float(celula.value).is_integer():
            return int(celula.value)
        return celula.value

    _limite(aba.nrows)
    linhas = [[valor(aba.cell(r, c)) for c in range(aba.ncols)] for r in range(aba.nrows)]
    livro.release_resources()
    while linhas and all(v == "" for v in linhas[-1]):
        linhas.pop()
    return linhas


def ler_docx(conteudo: bytes) -> Linhas:
    try:
        documento = Document(io.BytesIO(conteudo))
    except Exception as erro:
        raise ConversaoError("Documento Word inválido ou corrompido.") from erro
    if documento.tables:
        return [[c.text for c in linha.cells] for linha in documento.tables[0].rows]
    return [[p.text] for p in documento.paragraphs if p.text.strip()]


LEITORES = {
    "csv": ler_csv,
    "tsv": ler_csv,
    "json": ler_json,
    "yaml": ler_yaml,
    "yml": ler_yaml,
    "xml": ler_xml,
    "xlsx": ler_xlsx,
    "xls": ler_xls,
    "docx": ler_docx,
}


def ler_linhas(extensao: str, conteudo: bytes) -> Linhas:
    return [[limpar(c) for c in linha] for linha in LEITORES[extensao](conteudo)]


def ler_dados(extensao: str, conteudo: bytes) -> Any:
    """JSON, YAML e XML preservam a estrutura; os demais viram lista de registros."""
    if extensao == "json":
        return carregar_json(conteudo)
    if extensao in ("yaml", "yml"):
        return carregar_yaml(conteudo)
    if extensao == "xml":
        return carregar_xml(conteudo)
    return linhas_para_registros(ler_linhas(extensao, conteudo))


def serializavel(valor: Any) -> Any:
    """Converte datas e tipos não nativos para texto, aceitos por JSON, YAML e XML."""
    if isinstance(valor, dict):
        return {str(k): serializavel(v) for k, v in valor.items()}
    if isinstance(valor, (list, tuple, set)):
        return [serializavel(v) for v in valor]
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return limpar(valor)
    return str(valor)
