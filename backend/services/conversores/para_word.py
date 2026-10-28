import io
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, UnidentifiedImageError

from .erros import ConversaoError
from .tabular import Linhas, ler_linhas
from .texto import ler_blocos

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGENS = {"png", "jpg", "jpeg", "webp", "gif", "avif"}
TEXTOS = {"txt", "md", "html", "htm", "pptx"}
ORIGENS = {"pdf", "csv", "tsv", "xlsx", "xls", "json", "yaml", "yml", "xml"} | IMAGENS | TEXTOS


def _texto(extensao: str, conteudo: bytes) -> bytes:
    documento = Document()
    for tipo, texto in ler_blocos(extensao, conteudo):
        if tipo == "p":
            documento.add_paragraph(texto)
        else:
            documento.add_heading(texto, level=int(tipo[1]))
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def _pdf(conteudo: bytes) -> bytes:
    from pdf2docx import Converter

    with tempfile.TemporaryDirectory() as pasta:
        origem = Path(pasta) / "origem.pdf"
        destino = Path(pasta) / "destino.docx"
        origem.write_bytes(conteudo)
        conversor = Converter(str(origem))
        try:
            conversor.convert(str(destino))
        except Exception as erro:
            raise ConversaoError("Não foi possível converter este PDF para Word.") from erro
        finally:
            conversor.close()
        return destino.read_bytes()


def _tabela(linhas: Linhas) -> bytes:
    if not linhas:
        raise ConversaoError("O arquivo não possui dados para converter.")
    documento = Document()
    colunas = max(len(linha) for linha in linhas)
    tabela = documento.add_table(rows=0, cols=colunas)
    tabela.style = "Table Grid"
    for indice, linha in enumerate(linhas):
        celulas = tabela.add_row().cells
        for posicao, valor in enumerate(linha):
            celulas[posicao].text = str(valor)
            if indice == 0:
                for run in celulas[posicao].paragraphs[0].runs:
                    run.bold = True
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def _imagem(conteudo: bytes) -> bytes:
    try:
        imagem = Image.open(io.BytesIO(conteudo))
        imagem.load()
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as erro:
        raise ConversaoError("Imagem inválida ou formato não suportado.") from erro
    png = io.BytesIO()
    imagem.convert("RGBA" if imagem.mode in ("RGBA", "LA", "P") else "RGB").save(png, "PNG")
    png.seek(0)
    documento = Document()
    documento.add_picture(png, width=Inches(6))
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def converter(extensao: str, conteudo: bytes) -> tuple[bytes, str, str]:
    if extensao == "pdf":
        dados = _pdf(conteudo)
    elif extensao in IMAGENS:
        dados = _imagem(conteudo)
    elif extensao in TEXTOS:
        dados = _texto(extensao, conteudo)
    else:
        dados = _tabela(ler_linhas(extensao, conteudo))
    return dados, MIME, "docx"
