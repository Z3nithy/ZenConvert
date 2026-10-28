"""Extração de texto em blocos (título/parágrafo) para TXT, MD, HTML, DOCX, PDF e PPTX."""

import io
import re
from html.parser import HTMLParser

from .erros import ConversaoError
from .tabular import decodificar, limpar

Blocos = list[tuple[str, str]]
MAX_BLOCOS = 20_000
MAX_PAGINAS_PDF = 200
ORIGENS_TEXTO = {"txt", "md", "html", "htm", "docx", "pdf", "pptx"}


def _limitar(blocos: Blocos) -> Blocos:
    blocos = [(t, limpar(x).strip()) for t, x in blocos if limpar(x).strip()]
    if not blocos:
        raise ConversaoError("O arquivo não possui texto para converter.")
    if len(blocos) > MAX_BLOCOS:
        raise ConversaoError("O arquivo excede o limite de conteúdo permitido.", 413)
    return blocos


def _txt(conteudo: bytes) -> Blocos:
    texto = decodificar(conteudo).replace("\r\n", "\n")
    return [("p", p.replace("\n", " ")) for p in re.split(r"\n\s*\n", texto)]


def _md(conteudo: bytes) -> Blocos:
    blocos: Blocos = []
    paragrafo: list[str] = []

    def fechar():
        if paragrafo:
            blocos.append(("p", " ".join(paragrafo)))
            paragrafo.clear()

    for linha in decodificar(conteudo).splitlines():
        titulo = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", linha)
        if titulo:
            fechar()
            blocos.append((f"h{len(titulo.group(1))}", titulo.group(2)))
        elif not linha.strip():
            fechar()
        else:
            paragrafo.append(re.sub(r"^\s*(?:[-*+]|\d+\.)\s+", "• ", linha).strip())
    fechar()
    return blocos


class _Html(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocos: Blocos = []
        self._tipo = "p"
        self._partes: list[str] = []
        self._ignorar = 0

    def _fechar(self):
        texto = " ".join("".join(self._partes).split())
        if texto:
            self.blocos.append((self._tipo, texto))
        self._partes = []
        self._tipo = "p"

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self._ignorar += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._fechar()
            self._tipo = tag
        elif tag in ("p", "div", "li", "tr", "br", "section", "article", "blockquote"):
            self._fechar()

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self._ignorar = max(0, self._ignorar - 1)
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "div", "li", "tr", "td", "th"):
            self._fechar()

    def handle_data(self, dados):
        if not self._ignorar:
            self._partes.append(dados)


def _html(conteudo: bytes) -> Blocos:
    analisador = _Html()
    analisador.feed(decodificar(conteudo))
    analisador.close()
    analisador._fechar()
    return analisador.blocos


def _docx(conteudo: bytes) -> Blocos:
    from docx import Document

    try:
        documento = Document(io.BytesIO(conteudo))
    except Exception as erro:
        raise ConversaoError("Documento Word inválido ou corrompido.") from erro
    blocos: Blocos = []
    for item in documento.iter_inner_content():
        if hasattr(item, "rows"):
            for linha in item.rows:
                blocos.append(("p", " | ".join(c.text.strip() for c in linha.cells)))
            continue
        nome = (item.style.name if item.style is not None else "") or ""
        achou = re.match(r"(?:Heading|Título)\s*(\d)", nome)
        blocos.append((f"h{min(int(achou.group(1)), 6)}" if achou else "p", item.text))
    return blocos


def _pdf(conteudo: bytes) -> Blocos:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(conteudo)
    except Exception as erro:
        raise ConversaoError("PDF inválido, protegido ou corrompido.") from erro
    try:
        if len(pdf) > MAX_PAGINAS_PDF:
            raise ConversaoError(f"O PDF excede o limite de {MAX_PAGINAS_PDF} páginas.", 413)
        blocos: Blocos = []
        for pagina in pdf:
            texto = pagina.get_textpage().get_text_range().replace("\r\n", "\n")
            blocos += [("p", p.replace("\n", " ")) for p in re.split(r"\n\s*\n", texto)]
    finally:
        pdf.close()
    if not any(x.strip() for _, x in blocos):
        raise ConversaoError("O PDF não contém texto selecionável (pode ser digitalizado).")
    return blocos


def _pptx(conteudo: bytes) -> Blocos:
    from pptx import Presentation

    try:
        apresentacao = Presentation(io.BytesIO(conteudo))
    except Exception as erro:
        raise ConversaoError("Apresentação PPTX inválida ou corrompida.") from erro
    blocos: Blocos = []
    for numero, slide in enumerate(apresentacao.slides, start=1):
        titulo = slide.shapes.title
        blocos.append(("h2", (titulo.text_frame.text if titulo is not None else "") or f"Slide {numero}"))
        for forma in slide.shapes:
            if titulo is not None and forma.shape_id == titulo.shape_id:
                continue
            if forma.has_text_frame:
                blocos += [("p", p.text) for p in forma.text_frame.paragraphs]
            elif getattr(forma, "has_table", False) and forma.has_table:
                for linha in forma.table.rows:
                    blocos.append(("p", " | ".join(c.text.strip() for c in linha.cells)))
    return blocos


_LEITORES = {
    "txt": _txt,
    "md": _md,
    "html": _html,
    "htm": _html,
    "docx": _docx,
    "pdf": _pdf,
    "pptx": _pptx,
}


def ler_blocos(extensao: str, conteudo: bytes) -> Blocos:
    return _limitar(_LEITORES[extensao](conteudo))
